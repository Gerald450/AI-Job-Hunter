"""Extract plain text from uploaded resume files."""

from __future__ import annotations

import io
from pathlib import Path

from resumes.storage import ResumeValidationError


def extract_text(*, content: bytes, filename: str, content_type: str | None = None) -> str:
    """Return plain text from a PDF or DOCX resume."""
    ext = Path(filename or "").suffix.lower()
    ctype = (content_type or "").lower()

    if ext == ".pdf" or ctype == "application/pdf":
        return _extract_pdf(content)
    if ext == ".docx" or "wordprocessingml" in ctype:
        return _extract_docx(content)
    if ext == ".doc":
        raise ResumeValidationError(
            "Legacy .doc resumes are not supported for analysis. "
            "Upload a PDF or DOCX."
        )
    raise ResumeValidationError(
        "Unsupported resume type for text extraction. Upload a PDF or DOCX."
    )


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError, PdfStreamError
    except ImportError as exc:
        raise ResumeValidationError("pypdf is required to parse PDF resumes") from exc

    try:
        reader = PdfReader(io.BytesIO(content))
        parts: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    except (PdfReadError, PdfStreamError, OSError, ValueError) as exc:
        raise ResumeValidationError("Could not extract text from PDF resume") from exc

    joined = "\n".join(parts).strip()
    if not joined:
        raise ResumeValidationError("Could not extract text from PDF resume")
    return joined


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeValidationError(
            "python-docx is required to parse DOCX resumes"
        ) from exc

    document = Document(io.BytesIO(content))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # Also pull simple table cell text.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    parts.append(cell_text)
    joined = "\n".join(parts).strip()
    if not joined:
        raise ResumeValidationError("Could not extract text from DOCX resume")
    return joined
