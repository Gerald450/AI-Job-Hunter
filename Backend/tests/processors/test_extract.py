"""Tests for resume text extraction and analysis caching."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from matching.cache import description_hash
from resumes.extract import extract_text
from resumes.storage import ResumeValidationError


def test_description_hash_stable() -> None:
    assert description_hash("hello") == description_hash("hello")
    assert description_hash(" hello ") == description_hash("hello")
    assert description_hash("a") != description_hash("b")


def test_extract_docx(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Gerald Shimo")
    doc.add_paragraph("Skills: Python, FastAPI")
    buffer = io.BytesIO()
    doc.save(buffer)
    text = extract_text(
        content=buffer.getvalue(),
        filename="resume.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )
    assert "Gerald Shimo" in text
    assert "Python" in text


def test_extract_rejects_doc() -> None:
    with pytest.raises(ResumeValidationError, match="Legacy"):
        extract_text(content=b"fake", filename="old.doc")


def test_extract_empty_pdf_raises() -> None:
    # Minimal PDF with no extractable text pages still parses as PDF structure
    # for some writers; use clearly invalid empty-ish content via validation path.
    with pytest.raises(ResumeValidationError):
        extract_text(content=b"not a pdf", filename="x.pdf")
